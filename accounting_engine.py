# accounting_engine.py
# Modul Khusus Penanganan Kriptografi SHA-256 dan Manipulasi Berkas Excel

import os
import hashlib
import openpyxl
import shutil
from datetime import datetime
from openpyxl.styles import Font, PatternFill

class ImmutableAccountingEngine:
    def __init__(self, output_dir=r"C:\edge_ai\accounting_app_version"):
        self.output_dir = output_dir
        self.file_path = os.path.join(self.output_dir, "sistem_akuntansi_baku.xlsx")
        
        # Buat folder output utama dan folder cadangan jika belum ada
        self.backup_dir = os.path.join(self.output_dir, "backups")
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
        if not os.path.exists(self.backup_dir):
            os.makedirs(self.backup_dir)

    def hitung_hash_transaksi(self, tanggal, no_bukti, keterangan, ref, debit, kredit, hash_sebelumnya):
        """Membuat sidik jari kriptografi unik (SHA-256) untuk mengunci imutabilitas data."""
        blok_data = f"{tanggal}{no_bukti}{keterangan}{ref}{debit}{kredit}{hash_sebelumnya}"
        return hashlib.sha256(blok_data.encode('utf-8')).hexdigest()

    def inisialisasi_buku_akuntansi(self):
        """Membuat template file Excel akuntansi baru dengan struktur template yang benar."""
        wb = openpyxl.Workbook()
        
        ws_dashboard = wb.active
        ws_dashboard.title = "Dashboard"
        ws_jurnal_umum = wb.create_sheet(title="Jurnal Umum")
        wb.create_sheet(title="Buku Besar")
        
        ws_jurnal_umum.append(["Tanggal", "No. Bukti", "Keterangan", "Ref", "Debit", "Kredit", "Block Hash Token"])
        
        ws_dashboard.append(["DASHBOARD ANALISIS KEUANGAN KORPORAT"])
        ws_dashboard.append([])
        ws_dashboard.append(["Nama Akun Rekening", "Saldo Akhir Finansial"])
        
        wb.save(self.file_path)

    def dapatkan_hash_terakhir(self, ws_ju):
        """Membaca baris terakhir Jurnal Umum untuk mengambil rantai hash sebelumnya."""
        if ws_ju.max_row <= 1:
            return "0" * 64
        return ws_ju.cell(row=ws_ju.max_row, column=7).value

    def tambah_transaksi_imutabel(self, tanggal, no_bukti, keterangan, ref, debit, kredit):
        if not os.path.exists(self.file_path):
            self.inisialisasi_buku_akuntansi()
            
        wb = openpyxl.load_workbook(self.file_path)
        ws_ju = wb["Jurnal Umum"]
        
        hash_sebelumnya = self.dapatkan_hash_terakhir(ws_ju)
        hash_baru = self.hitung_hash_transaksi(tanggal, no_bukti, keterangan, ref, debit, kredit, hash_sebelumnya)
        
        ws_ju.append([tanggal, no_bukti, keterangan, ref, debit, kredit, hash_baru])
        
        self.proses_posting_buku_besar(wb)
        self.terapkan_desain_tabel(wb)
        
        wb.save(self.file_path)
        return True

    def proses_posting_buku_besar(self, wb):
        ws_ju = wb["Jurnal Umum"]
        
        if "Buku Besar" in wb.sheetnames:
            sheet_lama = wb["Buku Besar"]
            wb.remove(sheet_lama)
            
        ws_bb = wb.create_sheet(title="Buku Besar")
        ws_bb.append(["Nama Akun Rekening", "Kode", "Total Saldo Mutasi"])
        
        akun_map = {
            "101 - Kas Utama": {"debit": 0, "kredit": 0},
            "102 - Piutang Dagang": {"debit": 0, "kredit": 0},
            "201 - Utang Dagang": {"debit": 0, "kredit": 0},
            "301 - Modal Pemilik": {"debit": 0, "kredit": 0},
            "401 - Pendapatan Usaha": {"debit": 0, "kredit": 0},
            "501 - Beban Operasional": {"debit": 0, "kredit": 0}
        }
        
        for row in range(2, ws_ju.max_row + 1):
            ref_raw = ws_ju.cell(row=row, column=4).value
            ref_code = str(ref_raw).replace("[", "").replace("]", "").replace("'", "").replace('"', "").strip()
            
            deb_val = ws_ju.cell(row=row, column=5).value
            kre_val = ws_ju.cell(row=row, column=6).value
            
            try: deb_val = float(deb_val) if deb_val and deb_val != "-" else 0.0
            except: deb_val = 0.0
            try: kre_val = float(kre_val) if kre_val and kre_val != "-" else 0.0
            except: kre_val = 0.0
            
            for nama_akun in akun_map:
                if ref_code in nama_akun:
                    akun_map[nama_akun]["debit"] += deb_val
                    akun_map[nama_akun]["kredit"] += kre_val

        ws_dash = wb["Dashboard"]
        while ws_dash.max_row > 3:
            ws_dash.delete_rows(4)

        for akun, saldo in akun_map.items():
            kode = akun.split(" - ")[0]
            if kode in ["101", "102", "501"]:
                total_saldo = saldo["debit"] - saldo["kredit"]
            else:
                total_saldo = saldo["kredit"] - saldo["debit"]
            
            ws_bb.append([akun, kode, total_saldo])
            ws_dash.append([akun, total_saldo])

    def audit_rekonsiliasi_integritas(self):
        """Fitur Verifikasi Kriptografi + AUTO BACKUP saat verifikasi sukses."""
        if not os.path.exists(self.file_path):
            return False, "Berkas ledger tidak ditemukan."

        wb = openpyxl.load_workbook(self.file_path)
        ws_ju = wb["Jurnal Umum"]
        hash_sebelumnya = "0" * 64
        
        for row in range(2, ws_ju.max_row + 1):
            tanggal = str(ws_ju.cell(row=row, column=1).value or "").strip()
            no_bukti = str(ws_ju.cell(row=row, column=2).value or "").strip()
            keterangan = str(ws_ju.cell(row=row, column=3).value or "").strip()
            ref = str(ws_ju.cell(row=row, column=4).value or "").strip()
            
            debit_raw = ws_ju.cell(row=row, column=5).value
            kredit_raw = ws_ju.cell(row=row, column=6).value
            token_tercatat = ws_ju.cell(row=row, column=7).value

            debit = float(debit_raw) if debit_raw and debit_raw != "-" else 0.0
            kredit = float(kredit_raw) if kredit_raw and kredit_raw != "-" else 0.0
            
            if debit.is_integer(): debit = int(debit)
            if kredit.is_integer(): kredit = int(kredit)

            hash_kalkulasi = self.hitung_hash_transaksi(
                tanggal, no_bukti, keterangan, ref, debit, kredit, hash_sebelumnya
            )
            
            if hash_kalkulasi != token_tercatat:
                return False, f"MANIPULASI TERDETEKSI! Perubahan ilegal pada baris {row-1} ({no_bukti})."
            
            hash_sebelumnya = hash_kalkulasi
            
        # PROSES AUTO-BACKUP JIKA DATA VALID
        waktu_sekarang = datetime.now().strftime("%Y%m%dd_%H%M%S")
        nama_backup = f"backup_SAB_{waktu_sekarang}.xlsx"
        jalur_backup = os.path.join(self.backup_dir, nama_backup)
        shutil.copy2(self.file_path, jalur_backup)
        
        return True, f"Verifikasi sukses! Rantai data aman.\n[Auto-Backup Berhasil]: {nama_backup}"

    def terapkan_desain_tabel(self, wb):
        font_normal = Font(name="Segoe UI", size=11)
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1a1b26", end_color="1a1b26", fill_type="solid")
        
        for ws in wb.worksheets:
            if ws.views.sheetView:
                ws.views.sheetView[0].showGridLines = True
                
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.font = font_normal
                    if isinstance(cell.value, (int, float)) and cell.value > 0 and cell.row > 1:
                        cell.number_format = '"Rp"#,##0'
                    if cell.row == 1:
                        cell.font = font_header
                        cell.fill = fill_header
            
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 13)
